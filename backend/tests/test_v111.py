"""v1.11: idiomas de la interfaz, proyectos (eliminar/importar), extracción
robusta para Windows, caché de análisis y offsets UTF-16 (FR-072..076)."""

import io
import json
import zipfile

import docx

from backend import analysis, storage
from backend.extraction import extract_text
from backend.models import Document
from backend.nlp import kwic_occurrences, word_frequencies


def upload(client, name="e.txt", text="La escuela rural fue mi hogar.", headers=None):
    return client.post("/api/documents", files={"file": (name, text.encode("utf-8"), "text/plain")},
                       headers=headers or {})


# ----- Idioma de la interfaz (FR-072) -----

def test_errors_follow_interface_language(client):
    es = client.post("/api/codes", json={"domain": "Inventado", "name": "x"})
    en = client.post("/api/codes", json={"domain": "Inventado", "name": "x"},
                     headers={"X-Cua-Lang": "en-GB"})
    pt = client.post("/api/codes", json={"domain": "Inventado", "name": "x"},
                     headers={"X-Cua-Lang": "pt-BR"})
    assert "Categoría inválida" in es.json()["detail"]
    assert "Invalid category" in en.json()["detail"]
    assert "Categoria inválida" in pt.json()["detail"]


def test_validation_messages_are_translated_and_clean(client):
    r = client.post("/api/codes", json={"domain": "Emocional", "name": "  "},
                    headers={"X-Cua-Lang": "en"})
    assert r.status_code == 422
    assert r.json()["detail"] == "The code name cannot be empty"  # sin «Value error,»


def test_unknown_language_falls_back_to_spanish(client):
    r = client.get("/api/kwic?q=", headers={"X-Cua-Lang": "fr"})
    assert "Debes indicar" in r.json()["detail"]


def test_default_project_name_localized(client):
    project = client.get("/api/project", headers={"X-Cua-Lang": "pt"}).json()
    assert project["name"] == "Meu projeto"


def test_demo_in_each_language(client):
    for lang, name, quote in [
        ("es", "Proyecto de ejemplo (demo)", "miedo inicial"),
        ("en", "Example project (demo)", "initial fear"),
        ("pt", "Projeto de exemplo (demo)", "medo inicial"),
    ]:
        project = client.post("/api/projects/example", headers={"X-Cua-Lang": lang}).json()
        assert project["name"] == name
        code = project["codes"][0]
        text = project["documents"][0]["text"]
        assert code["name"] == quote
        assert text[code["start"]:code["end"]] == code["quote"]
        # Las claves de categoría no se traducen nunca
        assert {c["domain"] for c in project["codes"]} <= {
            "Emocional", "Descriptivo", "In Vivo", "Tensión/CHAT", "Proceso",
            "Teórico", "Relacional", "Crítico", "Método"}


def test_csv_headers_and_categories_localized(client):
    client.post("/api/codes", json={"domain": "Emocional", "name": "gozo"})
    es = client.get("/api/export.csv").content.decode("utf-8-sig")
    en = client.get("/api/export.csv?ui=en").content.decode("utf-8-sig")
    assert es.splitlines()[0] == "id,fecha,documento,dominio,codigo,cita,memo"
    assert en.splitlines()[0] == "id,date,document,category,code,quote,memo"
    assert ",Emotional,gozo," in en


def test_zip_readme_localized_and_keeps_project_json(client):
    upload(client)
    zf = zipfile.ZipFile(io.BytesIO(client.get("/api/export.zip?ui=pt").content))
    names = zf.namelist()
    assert "proyecto.json" in names and "LEIAME.txt" in names


def test_download_filename_survives_non_latin_names(client):
    client.patch(f"/api/projects/{client.get('/api/project').json()['id']}",
                 json={"name": "Educação 😀 ñandú"})
    r = client.get("/api/export.zip")
    assert r.status_code == 200
    assert "filename*=UTF-8''" in r.headers["content-disposition"]


def test_manual_per_language(client):
    for lang, marker in [("es", "Manual de uso"), ("en", "User manual"), ("pt", "Solução de problemas")]:
        r = client.get("/api/manual", headers={"X-Cua-Lang": lang})
        assert r.status_code == 200 and marker in r.json()["html"]


def test_manual_list_items_keep_their_continuation_lines(client):
    html = client.get("/api/manual").json()["html"]
    assert "</ul>\n<p>relaciones y exclusiones" not in html  # antes se partía el ítem
    assert "con relaciones y exclusiones) para explorar" in html


def test_custom_category_cannot_clone_translated_basic(client):
    r = client.post("/api/domains/custom", json={"name": "Emotional", "color": "#123456"})
    assert r.status_code == 409


# ----- Proyectos: eliminar e importar (FR-073) -----

def test_delete_project_requires_confirmation_and_goes_to_trash(client, data_dir):
    pid = client.post("/api/projects", json={"name": "Borrable"}).json()["id"]
    assert client.delete(f"/api/projects/{pid}").status_code == 400
    r = client.delete(f"/api/projects/{pid}?confirm=true")
    assert r.status_code == 200
    assert pid not in [p["id"] for p in client.get("/api/projects").json()]
    assert any(f.name.startswith(pid) for f in (data_dir / "papelera").iterdir())
    assert client.get("/api/project").json()["id"] != pid


def test_delete_unknown_project(client):
    assert client.delete("/api/projects/nope?confirm=true").status_code == 404


def test_import_zip_backup_roundtrip(client):
    upload(client, text="Texto con ñ y tildes: canción.")
    doc = client.get("/api/project").json()["documents"][0]
    client.post("/api/codes", json={"doc_id": doc["id"], "domain": "Proceso", "name": "c1",
                                    "quote": "Texto", "start": 0, "end": 5})
    original = client.get("/api/project").json()
    backup = client.get("/api/export.zip").content
    r = client.post("/api/projects/import",
                    files={"file": ("respaldo.zip", backup, "application/zip")})
    assert r.status_code == 201
    imported = r.json()
    assert imported["id"] != original["id"]  # identidad nueva: nunca sobrescribe
    assert imported["documents"][0]["text"] == "Texto con ñ y tildes: canción."
    assert imported["codes"][0]["name"] == "c1"
    assert imported["name"].endswith("(importado)")
    assert client.get("/api/project").json()["id"] == imported["id"]  # queda activo


def test_import_plain_json(client):
    project = client.post("/api/projects", json={"name": "Solo JSON"}).json()
    raw = json.dumps(project).encode("utf-8")
    r = client.post("/api/projects/import", files={"file": ("p.json", raw, "application/json")})
    assert r.status_code == 201


def test_import_rejects_garbage(client):
    r = client.post("/api/projects/import", files={"file": ("x.zip", b"PKnada", "application/zip")},
                    headers={"X-Cua-Lang": "en"})
    assert r.status_code == 422 and "not a valid Cua-libre project" in r.json()["detail"]


def test_snapshots_are_compressed_and_restorable(client):
    client.post("/api/codes", json={"domain": "Teórico", "name": "antes"})
    client.post("/api/project/save")
    snap = client.get("/api/snapshots").json()["snapshots"][0]["name"]
    assert snap.endswith(".json.gz")
    client.post("/api/codes", json={"domain": "Teórico", "name": "después"})
    client.post(f"/api/snapshots/{snap}/restore", json={"confirm": True})
    assert [c["name"] for c in client.get("/api/project").json()["codes"]] == ["antes"]


def test_failed_edit_leaves_no_half_applied_change(client):
    client.post("/api/codes", json={"domain": "Teórico", "name": "intacto"})
    r = client.patch("/api/codes/1", json={"name": "cambiado", "domain": "Falso"})
    assert r.status_code == 422
    assert client.get("/api/project").json()["codes"][0]["name"] == "intacto"


# ----- Extracción robusta (Windows) (FR-074) -----

def test_txt_with_bom_and_crlf():
    text = extract_text("win.txt", "﻿Línea uno\r\nLínea dos\r\n".encode("utf-8"))
    assert text == "Línea uno\nLínea dos\n"


def test_txt_windows_ansi_cp1252():
    raw = "“Comillas” y guión – acción".encode("cp1252")
    assert extract_text("ansi.txt", raw) == "“Comillas” y guión – acción"


def test_docx_tables_are_extracted_in_order():
    document = docx.Document()
    document.add_paragraph("Entrevista 1")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "E:"
    table.cell(0, 1).text = "¿Cómo llegaste a la escuela?"
    table.cell(1, 0).text = "P:"
    table.cell(1, 1).text = "Caminando dos horas."
    document.add_paragraph("Fin")
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("t.docx", buffer.getvalue())
    assert text.splitlines() == ["Entrevista 1", "E: | ¿Cómo llegaste a la escuela?",
                                 "P: | Caminando dos horas.", "Fin"]


# ----- Corpus: portugués, caché, offsets (FR-075/076) -----

def test_portuguese_stopwords():
    words = {w["word"] for w in word_frequencies(
        "A escola rural não era uma escola pobre, era outra forma de conhecer", lang="pt", min_len=2)}
    assert "escola" in words and "não" not in words and "uma" not in words


def test_nlp_portuguese_frequencies_and_grammatical_focus(client):
    # v1.12 (FR-077): el foco gramatical también existe para portugués
    upload(client, text="A escola rural era o coração da comunidade. A escola ensina e ensinava.")
    r = client.get("/api/nlp?lang=pt&min_len=4")
    assert r.status_code == 200 and r.json()["words"][0]["word"] == "escola"
    r = client.get("/api/nlp?lang=pt&pos=verb", headers={"X-Cua-Lang": "pt"})
    assert r.status_code == 200
    assert any(w["word"] == "ensinar" for w in r.json()["words"])  # lema de «ensina»


def test_kwic_offsets_are_utf16_after_emojis():
    text = "Mensaje 😀😀 de WhatsApp: la escuela cerró."
    doc = Document(id="d1", filename="w.txt", text=text)
    occ = kwic_occurrences([doc], "escuela", pos="all")[0]
    js_units = text.encode("utf-16-le")
    assert js_units[occ["start"] * 2:occ["end"] * 2].decode("utf-16-le") == "escuela"


def test_kwic_lemma_offsets_are_utf16_after_emojis():
    text = "😀 El niño cantó en la escuela."
    doc = Document(id="d1", filename="w.txt", text=text)
    occ = kwic_occurrences([doc], "cantar", lang="es", pos="verb")[0]
    js_units = text.encode("utf-16-le")
    assert js_units[occ["start"] * 2:occ["end"] * 2].decode("utf-16-le") == "cantó"


def test_offset_map_roundtrip():
    text = "a😀b😀c"
    offsets = analysis.OffsetMap(text)
    for cp in range(len(text) + 1):
        assert offsets.to_cp(offsets.to_utf16(cp)) == cp


def test_analysis_is_cached_on_disk_and_reused(client, monkeypatch):
    text = "La alegría de la escuela rural. El miedo del abandono."
    first = analysis.get(text, "es")
    assert analysis.is_ready(text, "es")
    analysis._mem.clear()  # simula reiniciar la app: queda la copia en disco

    def boom(*args, **kwargs):
        raise AssertionError("no debería recalcular")

    monkeypatch.setattr(analysis, "_compute", boom)
    assert analysis.get(text, "es") == first


def test_large_corpus_is_analysed_in_background(client, monkeypatch):
    monkeypatch.setattr(analysis, "SYNC_LIMIT_CHARS", 10)
    upload(client, text="Primer documento con alegría y esperanza compartida.")
    analysis._mem.clear()
    for f in storage.cache_dir().iterdir():
        f.unlink()
    r = client.get("/api/sentiment")
    if r.status_code == 202:  # el hilo de fondo puede terminar antes de consultar
        body = r.json()
        assert body["pending"] is True and body["total"] == 1
        import time

        for _ in range(100):
            if client.get("/api/sentiment").status_code == 200:
                break
            time.sleep(0.1)
    data = client.get("/api/sentiment").json()
    assert data["documents"][0]["matched"] >= 2
    status = client.get("/api/analysis/status").json()
    assert status == {"done": 1, "total": 1, "ready": True}


def test_long_single_document_beyond_spacy_limit_is_chunked(monkeypatch):
    monkeypatch.setattr(analysis, "CHUNK_CHARS", 60)
    text = "\n".join(["El niño cantó con alegría en la escuela rural."] * 8)
    data = analysis._compute(text, "es")
    verbs = [p for p in data["pos"] if p[3] == "cantar"]
    assert len(verbs) == 8
    for start, end, *_ in verbs:
        assert text[start:end] == "cantó"
