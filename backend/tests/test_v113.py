"""v1.13: importación de Word 97-2003 (.doc), y notas y cuadros de texto en
.docx (FR-082)."""

import io
import struct
import zipfile
from pathlib import Path

import docx
import olefile
import pytest
from docx.oxml import parse_xml

from backend import legacy_doc
from backend.extraction import ExtractionError, extract_text
from backend.tests.test_extraction import make_docx

# Entrevista ficticia en Word 97-2003 real (generada con textutil de macOS):
# título, un párrafo, una tabla hablante | intervención y dos párrafos finales
DOC = (Path(__file__).parent / "data" / "entrevista.doc").read_bytes()


# ----- .doc binario -----

def test_doc_word97():
    text = extract_text("entrevista.doc", DOC)
    lines = text.split("\n")
    assert lines[0] == "Entrevista 07 — Profesora de educación básica"
    assert "La niñez, la añoranza y la emoción pedagógica: «¿cómo se enseña" in lines[1]
    # las filas de la tabla, como en los .docx
    assert lines[2] == "Entrevistador | ¿Cómo recuerda su primer año en la escuela rural?"
    assert lines[3].startswith("Profesora | Fue difícil; llegué sin saber nada del lugar. “Aquí")
    assert lines[-1] == "Fin de la transcripción."
    assert not any(ord(c) < 32 and c not in "\n\t" for c in text)  # sin marcas de Word


def test_doc_through_the_api(client):
    r = client.post("/api/documents",
                    files={"file": ("Entrevista 07.DOC", DOC, "application/msword")})
    assert r.status_code == 201
    document = r.json()
    assert document["filename"] == "Entrevista 07.DOC"
    assert "pingüino del afiche" in document["text"]


def test_content_decides_not_the_extension():
    # un .docx guardado con extensión .doc, y un .doc antiguo renombrado a .docx
    assert "Vygotski" in extract_text("tesis.doc", make_docx())
    assert "pingüino" in extract_text("entrevista.docx", DOC)


def _patched_doc(tmp_path, offset: int, fmt: str, *values) -> bytes:
    """El fixture con algunos bytes de su cabecera cambiados."""
    path = tmp_path / "parche.doc"
    path.write_bytes(DOC)
    with olefile.OleFileIO(str(path), write_mode=True) as ole:
        stream = bytearray(ole.openstream("WordDocument").read())
        struct.pack_into(fmt, stream, offset, *values)
        ole.write_stream("WordDocument", bytes(stream))
    return path.read_bytes()


def test_doc_with_password(tmp_path):
    flags = struct.unpack_from("<H", _word_stream(DOC), 0x0A)[0]
    protected = _patched_doc(tmp_path, 0x0A, "<H", flags | 0x0100)
    with pytest.raises(ExtractionError, match="contraseña"):
        extract_text("secreto.doc", protected)


def test_doc_older_than_word97(tmp_path):
    word6 = _patched_doc(tmp_path, 0, "<HH", 0xA5DC, 101)
    with pytest.raises(ExtractionError, match="anterior a 1997"):
        extract_text("antiguo.doc", word6)


def _word_stream(data: bytes) -> bytes:
    with olefile.OleFileIO(io.BytesIO(data)) as ole:
        return ole.openstream("WordDocument").read()


def test_doc_unrecognised(client):
    with pytest.raises(ExtractionError, match="formato de Word reconocible"):
        extract_text("pagina.doc", b"<html><body>No soy un documento de Word</body></html>")
    r = client.post("/api/documents", headers={"X-Cua-Lang": "en"},
                    files={"file": ("page.doc", b"<html></html>", "application/msword")})
    assert r.status_code == 422 and "recognisable Word format" in r.json()["detail"]
    # el corpus queda intacto
    assert client.get("/api/project").json()["documents"] == []


def test_unsupported_message_names_doc():
    with pytest.raises(ExtractionError, match="PDF, DOCX, DOC o TXT"):
        extract_text("planilla.xls", b"\x00")


def test_piece_table_mixes_8bit_and_utf16():
    """Word reparte el texto en piezas: de un byte por carácter (con las
    comillas tipográficas de Windows-1252) o en UTF-16."""
    ansi = "“Año”…".encode("cp1252")   # 6 caracteres, en la posición 10
    wide = " niño €".encode("utf-16-le")  # 7 caracteres, en la posición 20
    stream = b"\x00" * 10 + ansi + b"\x00" * 4 + wide
    starts = struct.pack("<3I", 0, 6, 13)
    pieces = (struct.pack("<HIH", 0, 0x40000000 | (10 * 2), 0)
              + struct.pack("<HIH", 0, 20, 0))
    plc = starts + pieces
    clx = b"\x01" + struct.pack("<H", 2) + b"\xaa\xbb" + b"\x02" + struct.pack("<I", len(plc)) + plc
    assert legacy_doc._pieces_text(stream, clx, "x.doc") == "“Año”… niño €"


def test_fields_keep_only_the_visible_result():
    link = "Ver \x13 HYPERLINK \"https://example.org\" \x14la pauta\x15."
    assert legacy_doc._strip_fields(link) == "Ver la pauta."
    # cita de un gestor bibliográfico: la instrucción trae metadatos, no texto
    cite = "según \x13 ADDIN CSL_CITATION {\"id\": 1} \x14(Vygotski, 1934)\x15"
    assert legacy_doc._strip_fields(cite) == "según (Vygotski, 1934)"
    nested = "\x13 TOC \x14Capítulo 1\t\x13 PAGEREF _Toc1 \x143\x15\x15 fin"
    assert legacy_doc._strip_fields(nested) == "Capítulo 1\t3 fin"
    assert legacy_doc._strip_fields("a\x13 XE \"índice\" \x15b") == "ab"  # campo sin resultado


def test_layout_of_paragraphs_and_tables():
    raw = ("Título\r"
           "E\x07¿Pregunta?\x07\x07"
           "P\x07Línea uno\rLínea dos\x07\x07"
           "\x07celda sola\x07\x07"
           "salto\x0bmanual\x0cpágina nueva\x01\x08\r")
    assert legacy_doc._layout(raw).split("\n") == [
        "Título",
        "E | ¿Pregunta?",
        "P | Línea uno",
        "Línea dos",
        "celda sola",
        "salto", "manual", "página nueva",
    ]


# ----- RTF con extensión .doc -----

RTF = (r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\fswiss Helvetica;}}"
       r"{\colortbl;\red0\green0\blue0;}{\info{\title No debe aparecer}}" "\n"
       r"\pard\f0\fs24 Transcripci\'f3n de la ni\u241?ez\par" "\n"
       r"{\b Gl\u259?veanu} escribi\'f3 \ldblquote crear\rdblquote\~{\deleted borrado }\emdash  2019\par" "\n"
       r"{\field{\*\fldinst{HYPERLINK \"https://example.org\"}}{\fldrslt la pauta}}\par" "\n"
       r"\trowd\cellx3000\cellx9000\intbl Profesora\cell Aprend\'ed a escuchar\cell\row" "\n"
       r"\pard Llaves \{ y \} y barra \\ literales\par" "\n"
       r"{\pict\wmetafile8 0102030405}{\*\generator Programa de prueba;}}")


def test_rtf_saved_as_doc():
    text = extract_text("transcripcion.doc", RTF.encode("latin-1"))
    assert text.split("\n") == [
        "Transcripción de la niñez",
        "Glăveanu escribió “crear”\N{NO-BREAK SPACE}— 2019",
        "la pauta",
        "Profesora | Aprendí a escuchar",
        "Llaves { y } y barra \\ literales",
    ]


def test_rtf_skips_fallback_characters_and_binary_data():
    rtf = (r"{\rtf1\ansi\uc2 a\u960 pib{\uc0 \u241 o}\uc1 \u-10179?\u-8704?"
           r"{\pict\bin4 }{}{}{c}}").encode("latin-1")
    # \uc2: tras cada \u se saltan dos caracteres de respaldo; \bin4 salta los
    # 4 bytes siguientes («}{}{»), que son datos y no llaves
    assert legacy_doc.extract(rtf, "x.doc") == "aπbño\N{GRINNING FACE}c"


# ----- .docx: notas y cuadros de texto -----

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

TEXTBOX = f"""
<w:r xmlns:w="{W}" xmlns:mc="{MC}" xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps"><w:drawing><w:txbxContent>
      <w:p><w:r><w:t>Recuadro: recuerdos de la escuela</w:t></w:r></w:p>
    </w:txbxContent></w:drawing></mc:Choice>
    <mc:Fallback><w:pict><v:textbox><w:txbxContent>
      <w:p><w:r><w:t>Recuadro: recuerdos de la escuela</w:t></w:r></w:p>
    </w:txbxContent></v:textbox></w:pict></mc:Fallback>
  </mc:AlternateContent>
</w:r>"""

FOOTNOTES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W}">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t>Nota al pie: entrevista realizada en Ñuble.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""


def _docx_with_notes_and_textbox() -> bytes:
    document = docx.Document()
    document.add_paragraph("Cuerpo del documento.")
    document.paragraphs[0]._p.append(parse_xml(TEXTBOX))
    buffer = io.BytesIO()
    document.save(buffer)
    # python-docx no crea notas al pie: se agrega la parte a mano
    kind = "footnotes"
    result = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(result, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", (
                    f'<Override PartName="/word/{kind}.xml" ContentType="application/vnd.'
                    f'openxmlformats-officedocument.wordprocessingml.{kind}+xml"/></Types>'
                ).encode())
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", (
                    f'<Relationship Id="rIdNotas" Type="http://schemas.openxmlformats.org/'
                    f'officeDocument/2006/relationships/{kind}" Target="{kind}.xml"/></Relationships>'
                ).encode())
            target.writestr(item, data)
        target.writestr(f"word/{kind}.xml", FOOTNOTES)
    return result.getvalue()


def test_docx_notes_and_textboxes_go_at_the_end():
    text = extract_text("informe.docx", _docx_with_notes_and_textbox())
    assert text == ("Cuerpo del documento.\n\n"
                    "Nota al pie: entrevista realizada en Ñuble.\n\n"
                    "Recuadro: recuerdos de la escuela")  # una sola vez, no dos


def test_plain_docx_is_unchanged():
    assert extract_text("tesis.docx", make_docx("Solo un párrafo")) == "Solo un párrafo"


# ----- Versión -----

def test_version_is_the_same_everywhere():
    """La versión se escribe en tres lugares; los instaladores toman la de pyproject."""
    import tomllib

    from backend.app import APP_VERSION

    root = Path(__file__).resolve().parents[2]
    declared = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))["project"]["version"]
    assert APP_VERSION == declared
    assert f'<small id="version">v{declared}</small>' in (root / "frontend" / "index.html").read_text(encoding="utf-8")
